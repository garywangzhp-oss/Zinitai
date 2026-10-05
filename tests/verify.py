# -*- coding: utf-8 -*-
"""验证脚本：检查 scan 报告的类型分布 + apply 产物已真抹除。"""
import io
import re
import sys
from collections import Counter
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "tests" / "out"

print("== scan 报告类型分布 ==")
for name in ("示例采购合同_脱敏审查.md", "示例服务合同_脱敏审查.md"):
    t = io.open(OUT / name, encoding="utf-8").read()
    rows = re.findall(r"^\| \d+ \| ([^|]+) \|", t, re.M)
    print(f"  {name}: {dict(Counter(r.strip() for r in rows))}")
    for ln in t.splitlines():
        if "11010519900307891X" in ln or "44010619881122334X" in ln:
            cols = [c.strip() for c in ln.split("|")]
            print(f"    身份证 {cols[4]} -> 类型={cols[2]}（应为 身份证）")

print("== apply 产物真抹除检查 ==")
sensitive = [
    "13812345678", "021-65558899", "020-83334455", "13699998888",
    "11010519900307891X", "44010619881122334X",
    "91110108MA01B2C4XW", "91440101MA9Y2K7Q8R",
    "622202020011223349", "6225880137709987",
    "128,500.00", "350,000.00", "壹拾贰万捌仟伍佰元整", "叁拾伍万元整",
]
fail = False
docx_path = OUT / "示例采购合同_脱敏.docx"
pdf_path = OUT / "示例服务合同_脱敏.pdf"

from docx import Document  # noqa: E402
import pymupdf  # noqa: E402

def all_docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


dt = all_docx_text(docx_path)
pt = "\n".join(p.get_text() for p in pymupdf.open(pdf_path))
for tag, text in (("docx", dt), ("pdf", pt)):
    leaked = [s for s in sensitive if s in text]
    ph = len(re.findall(r"\[[^\]]+#\d+\]", text))
    print(f"  {tag}: 残留敏感原文 {len(leaked)} 处 {leaked or ''}，占位符 {ph} 处")
    if leaked:
        fail = True

print("== 订单号/编号未被误抹检查 ==")
if "1234567890123" not in dt:
    print("  ✗ docx 订单号 1234567890123 被误抹")
    fail = True
else:
    print("  ✓ docx 订单号保留")
if "JS-2026-0330" not in pt or "CG-2026-0818" not in dt:
    print("  ✗ 合同编号被误抹")
    fail = True
else:
    print("  ✓ 合同编号保留")

print("RESULT:", "FAIL" if fail else "PASS")
sys.exit(1 if fail else 0)
