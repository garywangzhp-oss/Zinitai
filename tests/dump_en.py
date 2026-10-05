# -*- coding: utf-8 -*-
"""查看英文 NDA 样本的全文，判断漏报/误报。"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
REAL = Path(__file__).resolve().parent / "real"

import pymupdf
from docx import Document

pdf = REAL / "en_Bonterms-Mutual-NDA-Version-1.pdf"
t = "\n".join(p.get_text() for p in pymupdf.open(pdf))
print(f"===== {pdf.name}（前 1800 字） =====")
print(t[:1800])

for name in ("en_Example-Cover-Page-for-Bonterms-Mutual-NDA .docx",
             "en_Playbook-of-Additional-Terms-for-Bonterms-NDA.docx"):
    doc = Document(REAL / name)
    parts = [p.text for p in doc.paragraphs]
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                parts.extend(p.text for p in c.paragraphs)
    txt = "\n".join(x for x in parts if x.strip())
    print(f"\n===== {name}（前 1500 字） =====")
    print(txt[:1500])
