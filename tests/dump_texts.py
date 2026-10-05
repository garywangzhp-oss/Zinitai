# -*- coding: utf-8 -*-
"""导出真实合同全文（限长），用于人工核对漏报误报。"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
REAL = Path(__file__).resolve().parent / "real"

import pymupdf
from docx import Document

pdf = REAL / "contract_harbin.pdf"
t = "\n".join(p.get_text() for p in pymupdf.open(pdf))
print(f"===== {pdf.name} 全文 =====")
print(t)

for name in ("offer-demo.docx", "jinan_ldht_6e96e579300e437381515a56aadb4cc7.docx",
             "henan_ldht_827940040e00437e980d61584b2a5e6b.docx"):
    p = REAL / name
    doc = Document(p)
    parts = [q.text for q in doc.paragraphs]
    for tb in doc.tables:
        for r in tb.rows:
            for c in r.cells:
                parts.extend(q.text for q in c.paragraphs)
    txt = "\n".join(x for x in parts if x.strip())
    print(f"\n===== {name} 全文（前 2500 字） =====")
    print(txt[:2500])
