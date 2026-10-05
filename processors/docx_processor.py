from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass

from docx import Document
from docx.text.paragraph import Paragraph

from detectors import scan_text
from detectors.base import Span
from .selection import filter_selected
from .selfcheck import residual_warnings


class DocxError(Exception):
    pass


@dataclass
class DocxHit:
    para: Paragraph
    desc: str            # 位置描述，如 "正文#12" / "表格[1,3]" / "页眉"
    spans: list[Span]


def _open(path):
    try:
        return Document(path)
    except Exception as exc:
        raise DocxError(f"无法打开 docx：{exc}（.doc 旧格式请先另存为 .docx）") from exc


def _iter_cell(cell, desc, out, seen):
    if id(cell._tc) in seen:
        return  # 合并单元格在 row.cells 里会重复出现，只处理一次
    seen.add(id(cell._tc))
    for p in cell.paragraphs:
        out.append((p, desc))
    for t in cell.tables:
        _iter_table(t, desc, out, seen)


def _iter_table(table, desc, out, seen):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            _iter_cell(cell, f"{desc}[{ri + 1},{ci + 1}]", out, seen)


def iter_paragraphs(doc):
    """正文段落 + 表格单元格段落（含嵌套表格）+ 页眉页脚段落。"""
    out: list[tuple[Paragraph, str]] = []
    for i, p in enumerate(doc.paragraphs):
        out.append((p, f"正文#{i + 1}"))
    seen: set[int] = set()
    for t in doc.tables:
        _iter_table(t, "表格", out, seen)
    for section in doc.sections:
        if section.header is not None:
            for p in section.header.paragraphs:
                out.append((p, "页眉"))
        if section.footer is not None:
            for p in section.footer.paragraphs:
                out.append((p, "页脚"))
    return out


def para_text(para: Paragraph) -> str:
    """run 拼接文本。扫描和替换都用它，保证偏移一致。

    注意：w:hyperlink 等容器里的文字不在 para.runs 中，会漏识别
    （合同正文极少用超链接，MVP 接受）。
    """
    return "".join(r.text or "" for r in para.runs)


def revision_warnings(doc) -> list[str]:
    """批注/修订里可能藏敏感信息，MVP 不处理它们，只警告。"""
    warns = []
    xml = doc.element.xml
    if "<w:ins " in xml or "<w:ins>" in xml:
        warns.append("文档含修订（插入痕迹），其中的文字未做脱敏，请人工检查")
    if "<w:del " in xml or "<w:delText" in xml:
        warns.append("文档含修订（删除痕迹），被删文字可能仍留在文件内，请人工检查")
    if "commentReference" in xml:
        warns.append("文档含批注，批注内容未做脱敏，请人工检查")
    return warns


def context_of(text: str, span: Span, width: int = 14) -> str:
    a = max(0, span.start - width)
    b = min(len(text), span.end + width)
    core = f"「{span.text}」"
    return f"…{text[a:span.start]}{core}{text[span.end:b]}…"


def _doc_text(doc) -> str:
    """整个文档（正文/表格/页眉页脚）拼接文本，用于抹后自检。"""
    return "\n".join(para_text(p) for p, _desc in iter_paragraphs(doc))


def scan_docx(path, use_ner: bool = False):
    """返回 (hits, warnings)。hit = (段落, 位置描述, spans)。"""
    doc = _open(path)
    hits: list[DocxHit] = []
    for para, desc in iter_paragraphs(doc):
        text = para_text(para)
        if not text.strip():
            continue
        spans = scan_text(text, use_ner)
        if spans:
            hits.append(DocxHit(para, desc, spans))
    return hits, revision_warnings(doc)


def replace_spans(para: Paragraph, spans: list[Span], assign) -> list[Span]:
    """跨 run 替换：命中的字符清掉，占位符落在命中起点所在 run。

    只重写内容有变化的 run，其余 run（含其中的软换行/制表符）原样保留。
    返回实际替换的 spans。
    """
    runs = para.runs
    if not runs or not spans:
        return []
    full = "".join(r.text or "" for r in runs)
    spans = [s for s in spans if 0 <= s.start < s.end <= len(full)
             and full[s.start:s.end] == s.text]
    if not spans:
        return []

    # 每个 run 的字符区间
    bounds, off = [], 0
    for r in runs:
        bounds.append((off, off + len(r.text or "")))
        off += len(r.text or "")

    replaced: list[Span] = []
    for i, r in enumerate(runs):
        s, e = bounds[i]
        txt = r.text or ""
        if e <= s:
            continue
        kept = []
        changed = False
        for j, ch in enumerate(txt):
            g = s + j
            hit = next((sp for sp in spans if sp.start <= g < sp.end), None)
            if hit is None:
                kept.append(ch)
            else:
                changed = True
                if g == hit.start:
                    kept.append(assign(hit))
                    replaced.append(hit)
        if changed:
            r.text = "".join(kept)
    # 去重（同一 span 只在其起点 run 记一次，防御性去重）
    seen, uniq = set(), []
    for sp in replaced:
        if id(sp) not in seen:
            seen.add(id(sp))
            uniq.append(sp)
    return uniq


def apply_docx(path_in, path_out, assign, exclude=frozenset(), use_ner=False, *, selected=None):
    """打开 → 逐段识别替换 → 另存。返回 (report_items, warnings, 替换处数)。

    exclude 为词库排除；selected（GUI 用，keyword-only）为 (label, text, 出现序次)
    命中集合——给定它时只抹集合内的 span，不给则抹除全部（减 exclude）。
    """
    doc = _open(path_in)
    warnings = list(revision_warnings(doc))
    items = []
    n = 0
    occ: dict = {}  # 出现序次按整个文件累计，与扫描端一致
    before_text = _doc_text(doc)
    redacted: Counter = Counter()
    for para, desc in iter_paragraphs(doc):
        text = para_text(para)
        if not text.strip():
            continue
        spans = [s for s in filter_selected(scan_text(text, use_ner), selected, occ)
                 if s.text not in exclude]
        if not spans:
            continue
        done = replace_spans(para, spans, assign)
        for sp in done:
            n += 1
            redacted[sp.text] += 1
            items.append({
                "placeholder": assign(sp),
                "label": sp.label,
                "original": sp.text,
                "context": context_of(text, sp),
                "location": desc,
                "note": sp.note,
            })
    # 抹后自检：回读文本，确认本次应抹的原文确实消失（自检失败不影响结果）
    try:
        warnings += residual_warnings(before_text, _doc_text(doc), dict(redacted))
    except Exception:
        pass
    doc.save(path_out)
    return items, warnings, n


def restore_docx(path_in, path_out, mapping: dict) -> int:
    """按映射表把占位符换回原文（docx）。返回替换处数。

    mapping: {占位符: 原文}。占位符含 closing bracket，互不为子串，不冲突。
    """
    doc = _open(path_in)
    n = 0
    for para, _desc in iter_paragraphs(doc):
        text = para_text(para)
        if not text.strip():
            continue
        spans = [Span(m.start(), m.end(), m.group(), "RESTORE")
                 for ph in mapping
                 for m in re.finditer(re.escape(ph), text)]
        if not spans:
            continue
        spans.sort(key=lambda s: s.start)
        n += len(replace_spans(para, spans, lambda s: mapping[s.text]))
    doc.save(path_out)
    return n
